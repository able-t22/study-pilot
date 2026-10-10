"""Multimodal ingestion: PDF / PPTX / video+audio / transcripts / text -> source-linked units."""
import contextvars, io, json, re, shutil, subprocess, time
from pathlib import Path
from . import config as C, db, llm
from .nlp import chunk_text, fmt_time

FAST_INGEST = contextvars.ContextVar("fast_ingest", default=False)
UNIT_BUFFER = contextvars.ContextVar("ingest_unit_buffer", default=None)
_ocr = None
VIDEO = {".mp4", ".mkv", ".webm", ".mov", ".avi", ".m4v"}
AUDIO = {".mp3", ".wav", ".m4a", ".flac", ".ogg"}


def ocr_image(img_bytes):
    global _ocr
    try:
        if _ocr is None:
            from rapidocr_onnxruntime import RapidOCR
            _ocr = RapidOCR()
        res, _ = _ocr(img_bytes)
        return " ".join(r[1] for r in (res or []))
    except Exception:
        return ""


def describe_image(img_bytes, context=""):
    """Figure -> text. Vision LLM if configured, else OCR of embedded labels."""
    if FAST_INGEST.get(): return ""  # readable text and scanned-page OCR still run
    if llm.vision_enabled():
        try:
            d = llm.chat("You describe figures from lecture material for a study knowledge base.",
                         "Describe this figure/diagram/slide image for a student in 2-5 sentences: what it shows, axis/labels/"
                         "components, relationships, and the key takeaway. Transcribe any equations. Nearby text: " + context[:400],
                         images=[img_bytes], max_tokens=500)
            if d and len(d) > 15: return "[Figure] " + d.strip()
        except Exception as e:
            print("[vision] failed:", e)
    t = ocr_image(img_bytes)
    return ("[Figure text via OCR] " + t) if len(t) > 15 else ""


def add_unit(sid, seq, kind, text, loc, label):
    row = (sid, seq, kind, text.strip(), json.dumps(loc), label)
    buffer = UNIT_BUFFER.get()
    if buffer is not None:
        buffer.append(row)
        if len(buffer) >= 100: _flush_units()
        return None
    return db.ex("INSERT INTO units(source_id,seq,kind,text,loc,label) VALUES(?,?,?,?,?,?)", row)


def _flush_units():
    buffer = UNIT_BUFFER.get()
    if buffer:
        db.many("INSERT INTO units(source_id,seq,kind,text,loc,label) VALUES(?,?,?,?,?,?)", buffer)
        buffer.clear()


def _save_media(sid, tag, data, ext="png"):
    fn = f"s{sid}_{tag}.{ext}"
    (C.MEDIA / fn).write_bytes(data)
    return fn


def ingest_pdf(sid, path, cb):
    import fitz
    doc = fitz.open(path); seq = 0; figs = 0
    for pno, page in enumerate(doc, 1):
        cb(f"PDF page {pno}/{len(doc)}", pno / len(doc))
        text = page.get_text("text")
        if len(text.strip()) < 30:  # scanned page -> OCR
            text = ocr_image(page.get_pixmap(dpi=130).tobytes("png"))
        for ch in chunk_text(text):
            seq += 1; add_unit(sid, seq, "text", ch, {"page": pno}, f"p.{pno}")
        if FAST_INGEST.get(): continue
        for img in page.get_images(full=True):
            if figs >= 60: break
            try:
                pix = fitz.Pixmap(doc, img[0])
                if pix.width < 120 or pix.height < 120: continue
                if pix.n - pix.alpha > 3: pix = fitz.Pixmap(fitz.csRGB, pix)
                data = pix.tobytes("png")
            except Exception:
                continue
            d = describe_image(data, text)
            if d:
                figs += 1; seq += 1
                _save_media(sid, f"p{pno}_{img[0]}", data)
                add_unit(sid, seq, "figure", d, {"page": pno}, f"p.{pno} (figure)")
        try:  # vector diagrams (no raster image)
            if not page.get_images() and len(page.get_drawings()) > 25 and figs < 60:
                d = describe_image(page.get_pixmap(dpi=110).tobytes("png"), text)
                if d:
                    figs += 1; seq += 1
                    add_unit(sid, seq, "figure", d, {"page": pno}, f"p.{pno} (diagram)")
        except Exception:
            pass
    return seq


def _shape_walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from _shape_walk(sh.shapes)
        else:
            yield sh


def ingest_pptx(sid, path, cb):
    from pptx import Presentation
    prs = Presentation(path); seq = 0
    pdf = None
    if not FAST_INGEST.get() and (shutil.which("soffice") or shutil.which("libreoffice")):  # exact slide rendering (optional)
        try:
            exe = shutil.which("soffice") or shutil.which("libreoffice")
            subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(C.MEDIA), str(path)], timeout=240,
                           capture_output=True)
            cand = C.MEDIA / (Path(path).stem + ".pdf")
            if cand.exists():
                pdf = str(cand.rename(C.MEDIA / f"s{sid}_slides.pdf"))
        except Exception as e:
            print("[pptx] soffice failed", e)
    if pdf: db.ex("UPDATE sources SET meta=? WHERE id=?", (json.dumps({"pdf": pdf}), sid))
    for i, slide in enumerate(prs.slides, 1):
        cb(f"Slide {i}/{len(prs.slides)}", i / len(prs.slides))
        title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        lines, imgs = [], []
        for sh in _shape_walk(slide.shapes):
            if sh.has_text_frame and sh != slide.shapes.title:
                lines += [p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip()]
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in r.cells))
            if sh.shape_type == 13:
                try: imgs.append(sh.image.blob)
                except Exception: pass
        lines = [l if l[-1] in ".?!:;" else l + "." for l in lines]
        body = ((title + ".\n") if title else "") + "\n".join(lines)
        if len(body.split()) >= 3:
            seq += 1; add_unit(sid, seq, "text", body, {"slide": i}, f"slide {i}")
        for k, data in enumerate(imgs):
            d = describe_image(data, body)
            if d:
                seq += 1; _save_media(sid, f"sl{i}_{k}", data)
                add_unit(sid, seq, "figure", d, {"slide": i}, f"slide {i} (figure)")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            seq += 1; add_unit(sid, seq, "notes", "Speaker notes: " + slide.notes_slide.notes_text_frame.text,
                               {"slide": i}, f"slide {i} (notes)")
    return seq


def parse_subs(path):
    """.vtt / .srt -> [(t0,t1,text)]"""
    raw = Path(path).read_text(encoding="utf-8", errors="ignore")
    out = []
    for blk in re.split(r"\n\s*\n", raw):
        m = re.search(r"(\d+):(\d+):(\d+)[\.,](\d+)\s*-->\s*(\d+):(\d+):(\d+)[\.,](\d+)", blk) or None
        if m:
            g = list(map(int, m.groups()))
            t0 = g[0] * 3600 + g[1] * 60 + g[2] + g[3] / 1000; t1 = g[4] * 3600 + g[5] * 60 + g[6] + g[7] / 1000
        else:
            m2 = re.search(r"(\d+):(\d+)[\.,](\d+)\s*-->\s*(\d+):(\d+)[\.,](\d+)", blk)
            if not m2: continue
            g = list(map(int, m2.groups())); t0 = g[0] * 60 + g[1] + g[2] / 1000; t1 = g[3] * 60 + g[4] + g[5] / 1000
        txt = " ".join(l for l in blk.splitlines() if "-->" not in l and not l.strip().isdigit() and l.strip() != "WEBVTT")
        txt = re.sub(r"<[^>]+>", "", txt).strip()
        if txt: out.append((t0, t1, txt))
    return out


def _window(segs, secs=55):
    cur, out = [], []
    for s in segs:
        if cur and s[1] - cur[0][0] > secs:
            out.append(cur); cur = []
        cur.append(s)
    if cur: out.append(cur)
    return out


_whisper = {}


def _whisper_model():
    """Load Whisper once and reuse it (loading takes many seconds and used to happen for every file)."""
    if C.WHISPER_MODEL not in _whisper:
        try: from faster_whisper import WhisperModel
        except ImportError: raise RuntimeError("Video/audio transcription is not installed on this deployment. Run StudyPilot locally (pip install -r requirements.txt) or upload a .vtt/.srt transcript, PDF or slides instead.")
        _whisper[C.WHISPER_MODEL] = WhisperModel(C.WHISPER_MODEL, device="cpu", compute_type="int8")
    return _whisper[C.WHISPER_MODEL]


def decode_audio(path, sr=16000):
    """Decode any audio/video file to 16 kHz mono float32 with PyAV.
    Done here (not inside faster-whisper) so a PyAV / faster-whisper version mismatch can no longer crash uploads
    (the old error: open() got an unexpected keyword argument 'metadata_errors'). Returns None if there is no audio track."""
    try: import av
    except ImportError: raise RuntimeError("Video/audio support is not installed. Run:  pip install av faster-whisper rapidocr-onnxruntime  (or pip install -r requirements.txt), then restart. Alternatively upload a .vtt/.srt transcript, PDF or slides.")
    import numpy as np
    chunks = []
    c = av.open(str(path))
    try:
        if not c.streams.audio: return None
        rs = av.audio.resampler.AudioResampler(format="s16", layout="mono", rate=sr)

        def take(frames):
            if frames is None: return
            for f in (frames if isinstance(frames, (list, tuple)) else [frames]):
                chunks.append(f.to_ndarray().reshape(-1))
        for fr in c.decode(audio=0):
            fr.pts = None
            take(rs.resample(fr))
        take(rs.resample(None))
    finally:
        c.close()
    return np.concatenate(chunks).astype("float32") / 32768.0 if chunks else None


def transcribe(path, cb):
    cb("Reading audio track...", 0.03)
    audio = decode_audio(path)
    if audio is None or len(audio) < 16000:
        cb("No speech track found - using on-screen content only", 0.1); return []
    dur = len(audio) / 16000
    cb("Loading speech model (first time downloads ~140 MB)...", 0.06)
    model = _whisper_model()
    segs, _info = model.transcribe(audio, vad_filter=True, beam_size=1); out = []
    for s in segs:  # generator: progress follows the audio position
        if s.text.strip(): out.append((s.start, s.end, s.text.strip()))
        cb(f"Transcribing {fmt_time(s.end)} of {fmt_time(dur)}", 0.08 + 0.52 * min(1.0, s.end / dur))
    return out


def ahash(img):
    s = img.convert("L").resize((8, 8)); px = list(s.getdata()); m = sum(px) / 64
    return [p > m for p in px]


def _duration(c):
    if c.duration: return c.duration / 1e6
    for st in c.streams.video:
        if st.duration and st.time_base: return float(st.duration * st.time_base)
    return 0.0


def keyframes(path, cb, max_frames=40):
    """Sample distinct frames spread across the WHOLE video (previously long videos were only covered up to ~30 min)."""
    try: import av
    except ImportError: raise RuntimeError("Video support is not installed. Run:  pip install av faster-whisper  then restart StudyPilot.")
    try:
        c = av.open(str(path))
        if not c.streams.video: c.close(); return
    except Exception:
        return
    try:
        dur = _duration(c); last = None; n = 0; t = 5.0
        step = max(float(C.FRAME_EVERY), dur / max_frames) if dur else float(C.FRAME_EVERY)
        while (t < dur or (not dur and n < 6)) and n < max_frames:
            cb(f"Sampling video frame at {fmt_time(t)}", 0.6 + 0.35 * min(1.0, t / max(dur, 1)))
            try:
                c.seek(int(t * 1e6)); frame = next(c.decode(video=0)); img = frame.to_image()
            except Exception:
                t += step; continue
            h = ahash(img)
            if last is None or sum(a != b for a, b in zip(h, last)) > 8:
                img.thumbnail((1024, 1024)); buf = io.BytesIO(); img.save(buf, "PNG")
                yield t, buf.getvalue(); last = h; n += 1
            t += step
    finally:
        c.close()


def ingest_media(sid, path, cb, is_video=True, subs=None):
    segs = parse_subs(subs) if subs else transcribe(path, cb)
    seq = 0
    for w in _window(segs):
        seq += 1
        add_unit(sid, seq, "transcript", " ".join(s[2] for s in w), {"t0": round(w[0][0], 1), "t1": round(w[-1][1], 1)},
                 f"{fmt_time(w[0][0])}-{fmt_time(w[-1][1])}")
    if is_video and not FAST_INGEST.get():
        _flush_units()
        tr = db.q("SELECT text, loc FROM units WHERE source_id=? AND kind='transcript'", (sid,))
        tr = [(json.loads(r["loc"]).get("t0", 0), r["text"]) for r in tr]
        for t, data in keyframes(path, cb):
            near = min(tr, key=lambda x: abs(x[0] - t))[1] if tr else ""
            d = describe_image(data, near)
            if d:
                seq += 1; _save_media(sid, f"t{int(t)}", data)
                add_unit(sid, seq, "figure", d, {"t0": round(t, 1), "t1": round(t + 5, 1)}, f"{fmt_time(t)} (on-screen)")
    return seq


def ingest_text(sid, path, cb):
    seq = 0
    for ch in chunk_text(Path(path).read_text(encoding="utf-8", errors="ignore")):
        seq += 1; add_unit(sid, seq, "text", ch, {"chunk": seq}, f"section {seq}")
    return seq


def ingest_file(path, name=None, cb=lambda m, p: None):
    path = Path(path); ext = path.suffix.lower(); name = name or path.name
    kind = {".pdf": "pdf", ".pptx": "slides", ".txt": "text", ".md": "text", ".vtt": "video", ".srt": "video"}.get(ext)
    if ext in VIDEO: kind = "video"
    if ext in AUDIO: kind = "audio"
    if not kind: raise ValueError(f"Unsupported file type: {ext}")
    meta = {}
    sid = db.ex("INSERT INTO sources(name,kind,path,meta,created) VALUES(?,?,?,?,?)", (name, kind, str(path), "{}", time.time()))
    token = UNIT_BUFFER.set([])
    try:
        if ext == ".pdf": n = ingest_pdf(sid, path, cb)
        elif ext == ".pptx": n = ingest_pptx(sid, path, cb)
        elif ext in VIDEO | AUDIO:
            sub = next((path.with_suffix(e) for e in (".vtt", ".srt") if path.with_suffix(e).exists()), None)
            n = ingest_media(sid, path, cb, ext in VIDEO, sub)
        elif ext in (".vtt", ".srt"):
            vid = next((path.with_suffix(e) for e in VIDEO if path.with_suffix(e).exists()), None)
            if vid:
                meta["media"] = str(vid)
            n = ingest_media(sid, path, cb, False, path)
        else: n = ingest_text(sid, path, cb)
        _flush_units()
    except Exception:
        db.ex("DELETE FROM units WHERE source_id=?", (sid,)); db.ex("DELETE FROM sources WHERE id=?", (sid,)); raise
    finally:
        UNIT_BUFFER.reset(token)
    if not meta.get("media") and ext in VIDEO | AUDIO: meta["media"] = str(path)
    old = json.loads(db.q1("SELECT meta FROM sources WHERE id=?", (sid,))["meta"] or "{}"); old.update(meta)
    db.ex("UPDATE sources SET meta=? WHERE id=?", (json.dumps(old), sid))
    if n == 0:
        db.ex("DELETE FROM sources WHERE id=?", (sid,)); raise ValueError("No extractable content found in " + name)
    return sid
